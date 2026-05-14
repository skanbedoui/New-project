from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = PROJECT_DIR / "outputs"
FIG_DIR = OUTPUT_DIR / "figures"
DOCX_PATH = PROJECT_DIR / "customer_churn_lab_report.docx"
FALLBACK_DOCX_PATH = PROJECT_DIR / "customer_churn_lab_report_logistic_regression.docx"

ACCENT = "2F6F73"
LIGHT_ACCENT = "EAF3F2"
LIGHT_GRAY = "F6F8FA"
BORDER = "D0D7DE"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(width_inches * 1440)))
    width.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_cell_text(cell, bold: bool = False, color: str | None = None, size: int = 9) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def add_table(
    doc: Document,
    df: pd.DataFrame,
    widths: list[float] | None = None,
    font_size: int = 9,
) -> None:
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])

    for idx, column in enumerate(df.columns):
        cell = table.rows[0].cells[idx]
        cell.text = str(column)
        set_cell_shading(cell, ACCENT)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        style_cell_text(cell, bold=True, color="FFFFFF", size=font_size)
        if widths:
            set_cell_width(cell, widths[idx])

    for row_index, (_, row) in enumerate(df.iterrows(), start=1):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.text = str(value)
            if row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            style_cell_text(cell, size=font_size)
            if widths:
                set_cell_width(cell, widths[idx])

    doc.add_paragraph()


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)


def add_figure(doc: Document, path: Path, caption: str, width: float = 5.9) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_ACCENT)
    set_cell_border(cell, color="B8D8D4", size="8")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    run.font.size = Pt(10)
    paragraph = cell.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    for run in paragraph.runs:
        run.font.size = Pt(9.5)
    doc.add_paragraph()


def set_document_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)

    for style_name, size in [("Heading 1", 15), ("Heading 2", 12)]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(ACCENT)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(ACCENT)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(90, 90, 90)


def add_footer(section) -> None:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "Customer Churn Prediction Lab Project"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(110, 110, 110)


def main() -> None:
    metrics = pd.read_csv(OUTPUT_DIR / "model_metrics.csv")
    tuning = pd.read_csv(OUTPUT_DIR / "tuning_summary.csv")
    importance = pd.read_csv(OUTPUT_DIR / "feature_importance.csv").head(10)
    examples = pd.read_csv(OUTPUT_DIR / "sample_predictions.csv").head(5)
    profile = json.loads((OUTPUT_DIR / "data_profile.json").read_text(encoding="utf-8"))

    for col in ["Accuracy", "Precision", "Recall", "F1", "ROC-AUC"]:
        metrics[col] = metrics[col].map(lambda value: f"{value:.4f}")
    importance["Importance"] = importance["Importance"].map(lambda value: f"{value:.4f}")
    examples = examples[["Actual Churn", "Predicted Churn", "Churn Probability"]].copy()
    examples["Churn Probability"] = examples["Churn Probability"].map(lambda value: f"{value:.2%}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    add_footer(section)
    set_document_styles(doc)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Customer Churn Prediction")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Business / Marketing / Supply Chain Lab Project Report")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(75, 85, 99)

    add_callout(
        doc,
        "Project overview",
        "This report summarizes a supervised machine learning project for predicting customer churn. "
        "The project compares Random Forest, Logistic Regression, and SVM using one shared preprocessing pipeline, "
        "cross-validation based hyperparameter tuning, and a common unseen test set.",
    )

    summary = pd.DataFrame(
        {
            "Item": ["Dataset rows", "Modeling sample", "Selected model", "Domain", "Target"],
            "Value": [
                f"{profile['clean_rows']:,}",
                f"{profile['sample_rows']:,}",
                profile["best_model"],
                "Business / Marketing / Supply Chain",
                "Churn",
            ],
        }
    )
    add_table(doc, summary, widths=[1.7, 4.5], font_size=9)

    doc.add_page_break()

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Customer churn is a central business problem because losing active customers reduces revenue "
        "and increases the cost of growth. In marketing and customer success teams, a churn model can "
        "support retention campaigns by ranking customers according to churn risk."
    )
    doc.add_paragraph(
        "The objective is to predict the binary target Churn, where 1 means the customer churned and "
        "0 means the customer stayed. This is a supervised classification task with measurable outcomes."
    )
    add_figure(
        doc,
        FIG_DIR / "target_distribution.png",
        "Figure 1. Churn target distribution.",
        width=4.8,
    )

    doc.add_heading("2. Background", level=1)
    doc.add_paragraph(
        "The dataset contains demographic, behavioral, contract, spending, and interaction variables. "
        "These variables are relevant because churn is often influenced by customer age, tenure, product "
        "usage, support burden, payment delay, contract length, subscription type, spending, and recency."
    )
    feature_df = pd.DataFrame(
        {
            "Feature group": ["Numeric", "Categorical", "Removed identifier", "Target"],
            "Columns": [
                ", ".join(profile["numeric_features"]),
                ", ".join(profile["categorical_features"]),
                "CustomerID",
                "Churn",
            ],
        }
    )
    add_table(doc, feature_df, widths=[1.7, 4.7], font_size=8)

    doc.add_page_break()

    doc.add_heading("3. Methodology", level=1)
    doc.add_paragraph(
        "Rows with missing values were removed because the file only contained one incomplete row. "
        "CustomerID was excluded because it is an identifier rather than a meaningful predictor. Numeric "
        "features were median-imputed and scaled. Categorical features were mode-imputed and one-hot encoded."
    )
    doc.add_paragraph(
        "A stratified sample was used for modeling so that the notebook remains practical to rerun while "
        "preserving the target distribution. The sampled data was split into 80 percent training and "
        "20 percent testing. All three models used the same split and preprocessing pipeline."
    )
    doc.add_heading("Hyperparameter Tuning Summary", level=2)
    add_table(tuning_doc := doc, tuning, widths=[1.2, 1.1, 4.4], font_size=7)

    doc.add_heading("4. Algorithms", level=1)
    doc.add_paragraph(
        "Random Forest is an ensemble of decision trees. It reduces overfitting by averaging many trees "
        "trained on randomized subsets of the data and features. Important hyperparameters include the "
        "number of trees and maximum tree depth."
    )
    doc.add_paragraph(
        "Logistic Regression is a linear classification model. It estimates the probability of churn by "
        "combining feature values with learned coefficients. Important hyperparameters include the "
        "regularization strength C and the class weighting strategy."
    )
    doc.add_paragraph(
        "SVM searches for a separating boundary with a maximum margin. A linear SVM was used for speed "
        "and interpretability on the large tabular dataset, with probability calibration for fair ROC-AUC "
        "comparison."
    )

    doc.add_page_break()

    doc.add_heading("5. Results and Interpretation", level=1)
    doc.add_paragraph(
        "The comparison uses accuracy, precision, recall, F1-score, and ROC-AUC. F1-score is emphasized "
        "because churn prediction requires a balance between finding true churners and limiting false "
        "retention alarms."
    )
    add_table(doc, metrics, widths=[1.3, 0.85, 0.9, 0.75, 0.65, 0.85], font_size=8)
    add_figure(
        doc,
        FIG_DIR / "model_comparison.png",
        "Figure 2. Model comparison across evaluation metrics.",
        width=6.1,
    )

    doc.add_heading("6. Model Selection and Inference", level=1)
    doc.add_paragraph(
        f"The selected model is {profile['best_model']} because it achieved the strongest F1-score on "
        "the unseen test data. The ROC and confusion matrix plots support this choice by showing strong "
        "class separation and a low number of mistakes."
    )
    add_figure(doc, FIG_DIR / "roc_curves.png", "Figure 3. ROC curve comparison.", width=5.4)
    add_figure(
        doc,
        FIG_DIR / "confusion_matrices.png",
        "Figure 4. Confusion matrices on the unseen test set.",
        width=6.0,
    )

    doc.add_page_break()

    doc.add_heading("7. Explainability and Conclusion", level=1)
    doc.add_paragraph(
        "The top feature chart helps explain which customer attributes influenced the selected model most. "
        "This is useful for a business audience because it connects prediction quality to actionable churn "
        "drivers."
    )
    add_figure(
        doc,
        FIG_DIR / "feature_importance.png",
        f"Figure 5. Top features used by {profile['best_model']}.",
        width=5.7,
    )
    doc.add_heading("Top Feature Table", level=2)
    add_table(doc, importance, widths=[4.8, 1.2], font_size=8)

    doc.add_page_break()
    doc.add_heading("Example Inference Output", level=2)
    add_table(doc, examples, widths=[1.5, 1.6, 1.6], font_size=8)

    doc.add_paragraph(
        "In conclusion, the project satisfies the core machine learning workflow: clear objective, dataset "
        "preparation, preprocessing, model tuning, fair comparison, final model selection, and inference "
        "on unseen data. The notebook provides the executable workflow, while the Streamlit dashboard "
        "presents the results interactively."
    )

    try:
        doc.save(DOCX_PATH)
        print(f"Created {DOCX_PATH}")
    except PermissionError:
        doc.save(FALLBACK_DOCX_PATH)
        print(f"Created {FALLBACK_DOCX_PATH}")


if __name__ == "__main__":
    main()
